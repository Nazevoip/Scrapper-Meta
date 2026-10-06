"""Blocos de desenho do Kit Massoterapia Organizada em PPTX editavel (A4).

Tudo em pontos (1 pt = 1/72"). A pagina A4 tem 595,28 x 841,89 pt.
"""

from __future__ import annotations

import os
from functools import lru_cache
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

W, H = 595.28, 841.89
ASSETS = Path(__file__).with_name("assets")
FONTES = Path(os.path.expanduser("~/.fonts"))

ARQ_FONTE = {
    ("Montserrat", False): "Montserrat-Regular-static.ttf",
    ("Montserrat", True): "Montserrat-Bold-static.ttf",
    ("Cormorant Garamond", False): "CormorantGaramond-SemiBold-static.ttf",
    ("Cormorant Garamond", True): "CormorantGaramond-Bold-static.ttf",
    ("Libre Baskerville", False): "LibreBaskerville-Regular-static.ttf",
    ("Libre Baskerville", True): "LibreBaskerville-Bold-static.ttf",
}


@lru_cache(None)
def _fonte(nome: str, negrito: bool) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONTES / ARQ_FONTE[(nome, negrito)]), 100)


def largura(txt: str, fonte: str, tam: float, negrito: bool = False, espaco: float = 0) -> float:
    return _fonte(fonte, negrito).getlength(txt) * tam / 100 + espaco * max(len(txt) - 1, 0)


def quebrar(txt: str, fonte: str, tam: float, larg: float, negrito: bool = False) -> int:
    """Quantas linhas o texto ocupa nessa largura."""
    total = 0
    for par in txt.split("\n"):
        atual, n = "", 1
        for palavra in par.split():
            tenta = f"{atual} {palavra}".strip()
            if largura(tenta, fonte, tam, negrito) <= larg:
                atual = tenta
            else:
                n += 1
                atual = palavra
        total += n
    return total


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


class Doc:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(W), Pt(H)
        self.layout = self.prs.slide_layouts[6]

    def pagina(self, fundo: str) -> "Pag":
        s = self.prs.slides.add_slide(self.layout)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(fundo)
        return Pag(s)

    def salvar(self, caminho):
        self.prs.save(caminho)


class Pag:
    def __init__(self, slide):
        self.slide = slide
        self.sh = slide.shapes

    def _atras(self, *shapes):
        """Manda formas para tras (logo apos o cabecalho do spTree)."""
        arvore = self.sh._spTree
        for i, s in enumerate(shapes):
            arvore.remove(s._element)
            arvore.insert(2 + i, s._element)

    def caixa(self, x, y, w, h, fundo=None, borda=None, esp=0.75, raio=None, forma=None):
        tipo = forma or (MSO_SHAPE.ROUNDED_RECTANGLE if raio else MSO_SHAPE.RECTANGLE)
        s = self.sh.add_shape(tipo, Pt(x), Pt(y), Pt(w), Pt(h))
        if raio and tipo in (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.ROUND_2_SAME_RECTANGLE):
            s.adjustments[0] = min(0.5, raio / min(w, h))
        if fundo:
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(fundo)
        else:
            s.fill.background()
        if borda:
            s.line.color.rgb = rgb(borda)
            s.line.width = Pt(esp)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        _sem_estilo(s)
        s.text_frame.text = ""
        return s

    def oval(self, x, y, w, h, fundo=None, borda=None, esp=0.75):
        return self.caixa(x, y, w, h, fundo, borda, esp, forma=MSO_SHAPE.OVAL)

    def texto(self, x, y, w, h, txt, fonte, tam, cor, negrito=False, italico=False,
              alinha="l", ancora="t", espaco=0, entrelinha=None, apos=0):
        tb = self.sh.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[ancora]
        for i, par in enumerate(txt.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[alinha]
            if entrelinha:
                p.line_spacing = entrelinha
            if apos:
                p.space_after = Pt(apos)
            r = p.add_run()
            r.text = par
            f = r.font
            f.name, f.size, f.bold, f.italic = fonte, Pt(tam), negrito, italico
            f.color.rgb = rgb(cor)
            if espaco:
                r._r.get_or_add_rPr().set("spc", str(int(espaco * 100)))
        return tb

    def linha(self, x1, y1, x2, y2, cor, esp=0.6):
        c = self.sh.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        c.line.color.rgb = rgb(cor)
        c.line.width = Pt(esp)
        _sem_estilo(c)
        return c

    def imagem(self, nome, x, y, w=None, h=None):
        return self.sh.add_picture(str(ASSETS / nome), Pt(x), Pt(y),
                                   Pt(w) if w else None, Pt(h) if h else None)

    def arco(self, x, y, w, h, cor, esp=0.8, inicio=180, fim=0):
        s = self.sh.add_shape(MSO_SHAPE.ARC, Pt(x), Pt(y), Pt(w), Pt(h))
        s.adjustments[0], s.adjustments[1] = inicio, fim
        s.fill.background()
        s.line.color.rgb = rgb(cor)
        s.line.width = Pt(esp)
        s.shadow.inherit = False
        _sem_estilo(s)
        return s

    def tabela(self, x, y, larguras, alt_linha, cabecalho, n_linhas, estilo, primeira_coluna=None):
        """Tabela editavel. `estilo` traz cores e fontes do cabecalho e das linhas."""
        cols = len(larguras)
        gf = self.sh.add_table(n_linhas + 1, cols, Pt(x), Pt(y), Pt(sum(larguras)),
                               Pt(alt_linha * (n_linhas + 1)))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tblPr.set(attr, "0")
        for j, w in enumerate(larguras):
            tbl.columns[j].width = Pt(w)
        for i in range(n_linhas + 1):
            tbl.rows[i].height = Pt(estilo.get("alt_cab", alt_linha) if i == 0 else alt_linha)
            for j in range(cols):
                cel = tbl.cell(i, j)
                cel.margin_left = cel.margin_right = Pt(4)
                cel.margin_top = cel.margin_bottom = Pt(1)
                cel.vertical_anchor = MSO_ANCHOR.MIDDLE
                if i == 0:
                    fundo, txt = estilo["fundo_cab"], cabecalho[j]
                else:
                    fundo = estilo["fundo"][(i - 1) % len(estilo["fundo"])]
                    txt = ""
                    if primeira_coluna and j == 0:
                        txt = primeira_coluna[i - 1]
                cel.fill.solid()
                cel.fill.fore_color.rgb = rgb(fundo)
                tf = cel.text_frame
                tf.text = txt
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER if (i == 0 and estilo.get("cab_centro", True)) else (
                    PP_ALIGN.CENTER if estilo.get("corpo_centro") else PP_ALIGN.LEFT)
                # celula vazia tambem precisa do tamanho de fonte, senao a linha fica alta
                tam_cel = estilo["tam_cab"] if i == 0 else estilo["tam"]
                p._p.get_or_add_endParaRPr().set("sz", str(int(tam_cel * 100)))
                for r in p.runs:
                    f = r.font
                    if i == 0:
                        f.name, f.size, f.bold = estilo["fonte_cab"], Pt(estilo["tam_cab"]), estilo.get("negrito_cab", False)
                        f.color.rgb = rgb(estilo["cor_cab"])
                    else:
                        f.name, f.size = estilo["fonte"], Pt(estilo["tam"])
                        f.color.rgb = rgb(estilo["cor"])
                _bordas(cel, estilo["borda"], estilo.get("esp_borda", 0.5))
        return gf


def _bordas(cel, cor, esp):
    tcPr = cel._tc.get_or_add_tcPr()
    # No XML as bordas (lnL, lnR, lnT, lnB) vem antes do preenchimento da celula.
    for k, lado in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        velho = tcPr.find(qn(lado))
        if velho is not None:
            tcPr.remove(velho)
        ln = tcPr.makeelement(qn(lado), {"w": str(int(esp * 12700)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        solid = ln.makeelement(qn("a:solidFill"), {})
        solid.append(solid.makeelement(qn("a:srgbClr"), {"val": cor.lstrip("#").upper()}))
        ln.append(solid)
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
        tcPr.insert(k, ln)


def _sem_estilo(shape):
    """Remove o <p:style> do tema: sem sombra nem efeitos herdados (cores ja sao explicitas)."""
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)
